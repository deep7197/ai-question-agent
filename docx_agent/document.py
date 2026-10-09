"""
A Word document, flattened into the blocks an answer can go into.

Unlike Google Docs - where the page is painted onto a canvas and an
extension can only aim the cursor and hope - a .docx can be edited
exactly: the empty cell under a question is an object, and its text
is set directly. Nothing is guessed, and nothing already written is
touched.

Blocks come out in reading order, paragraphs and table cells alike,
and empty ones are kept: an empty cell or an empty paragraph is the
space the document leaves for an answer.
"""

from dataclasses import dataclass, field
from typing import List, Optional

import docx
from docx.oxml import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


@dataclass
class Block:
    """One paragraph, or one cell of a table."""

    index: int
    text: str
    in_table: bool

    paragraph: Optional[Paragraph] = None
    cell: Optional[_Cell] = None

    table_at: int = -1
    row: int = -1
    column: int = -1

    @property
    def empty(self) -> bool:
        return not self.text.strip()


@dataclass
class Grid:
    """A table a question expects to be filled in."""

    table_at: int
    headers: List[str]
    rows: List[str]
    columns: int = 0
    table: Optional[Table] = None


@dataclass
class Document:
    """The whole file, and the blocks it is made of."""

    path: str
    blocks: List[Block] = field(default_factory=list)
    tables: List[Table] = field(default_factory=list)
    _doc: docx.Document = None

    def save(self, path: str = None) -> str:
        self._doc.save(path or self.path)
        return path or self.path


def _cell_text(cell: _Cell) -> str:
    parts = [p.text.strip() for p in cell.paragraphs]
    return "\n".join(part for part in parts if part).strip()


def read(path: str) -> Document:
    """Loads a .docx and flattens it into blocks, in reading order."""

    document = docx.Document(path)

    blocks: List[Block] = []
    tables: List[Table] = []

    def add(block: Block) -> None:
        block.index = len(blocks)
        blocks.append(block)

    def walk_table(table: Table) -> None:

        table_at = len(tables)
        tables.append(table)

        """
        A cell merged across rows or columns is handed back once for
        every row and column it covers. Assessment workbooks merge
        constantly - a question cell spanning the rows beside
        "Satisfactory response / Yes / No" - so without this every
        such question would be found two or three times over.
        """
        seen = set()

        for row_at, row in enumerate(table.rows):
            for column_at, cell in enumerate(row.cells):

                marker = id(cell._tc)

                if marker in seen:
                    continue

                seen.add(marker)

                text = _cell_text(cell)

                nested = cell.tables

                if text or not nested:
                    add(Block(
                        index=0,
                        text=text,
                        in_table=True,
                        cell=cell,
                        table_at=table_at,
                        row=row_at,
                        column=column_at,
                    ))

                for inner in nested:
                    walk_table(inner)

    for child in document.element.body.iterchildren():

        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            add(Block(
                index=0,
                text=paragraph.text.strip(),
                in_table=False,
                paragraph=paragraph,
            ))

        elif child.tag.endswith("}tbl"):
            walk_table(Table(child, document))

    return Document(path=path, blocks=blocks, tables=tables, _doc=document)


def style_from(block: Block):
    """The run formatting of an existing block, to match when writing."""

    paragraph = block.paragraph

    if paragraph is None and block.cell is not None:
        paragraphs = block.cell.paragraphs
        paragraph = paragraphs[0] if paragraphs else None

    if paragraph is None or not paragraph.runs:
        return None

    run = paragraph.runs[0]

    return {
        "name": run.font.name,
        "size": run.font.size,
        "bold": False,
        "italic": False,
    }


def _apply(paragraph: Paragraph, text: str, style) -> None:

    run = paragraph.add_run(text)

    if style:
        run.font.name = style.get("name")
        if style.get("size"):
            run.font.size = style["size"]
        run.bold = False
        run.italic = False


def fill(block: Block, text: str, style=None) -> None:
    """
    Writes an answer into an empty block. Paragraph breaks in the
    answer become separate paragraphs, the way a person would type
    them.
    """

    paragraphs = [
        " ".join(line.strip() for line in part.splitlines() if line.strip())
        for part in text.split("\n\n")
    ]

    paragraphs = [part for part in paragraphs if part] or [""]

    if block.cell is not None:

        target = block.cell

        first = target.paragraphs[0] if target.paragraphs \
            else target.add_paragraph()

        _apply(first, paragraphs[0], style)

        for extra in paragraphs[1:]:
            _apply(target.add_paragraph(), extra, style)

        return

    _apply(block.paragraph, paragraphs[0], style)

    anchor = block.paragraph

    for extra in paragraphs[1:]:
        anchor = _insert_after(anchor, extra, style)


def _insert_after(paragraph: Paragraph, text: str, style) -> Paragraph:

    new_element = paragraph._p.makeelement(paragraph._p.tag, {})
    paragraph._p.addnext(new_element)

    created = Paragraph(new_element, paragraph._parent)
    created.style = paragraph.style

    _apply(created, text, style)

    return created


def insert_after(block: Block, text: str, style=None) -> None:
    """
    For a question the document leaves no room under: the answer is
    added directly beneath it, rather than anywhere else.

    When the question sits in a box, the answer goes below that box
    as its own paragraph - not appended inside the same cell, where
    it would read as another line of the question.
    """

    paragraphs = [
        " ".join(line.strip() for line in part.splitlines() if line.strip())
        for part in text.split("\n\n")
    ]

    paragraphs = [part for part in paragraphs if part] or [""]

    if block.cell is not None:

        anchor_element = block.cell._tc.getparent().getparent()

        for part in paragraphs:

            new_element = OxmlElement("w:p")
            anchor_element.addnext(new_element)
            anchor_element = new_element

            _apply(
                Paragraph(new_element, block.cell._parent),
                part,
                style,
            )

        return

    anchor = block.paragraph

    for part in paragraphs:
        anchor = _insert_after(anchor, part, style)
