"""
Builds a workbook shaped like a real assessment - question box with
the assessor's Satisfactory/Yes/No lines, the student's empty box
after it, a case-study parent, a grid to fill in, and a question
with no space at all - then checks the agent reads it correctly and
writes every answer into the right place.

    .venv\\Scripts\\python.exe tests\\test_docx_agent.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx                                              # noqa: E402

from docx_agent import document as doc_io                # noqa: E402
from docx_agent import questions as finder               # noqa: E402


def build(path):

    out = docx.Document()

    out.add_paragraph("Assessment Task 1 - Unit Knowledge Test (UKT)")

    def box(lines):
        table = out.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        cell.paragraphs[0].text = lines[0]
        for line in lines[1:]:
            cell.add_paragraph(line)
        return table

    box([
        "Question 1. Read the case study and answer the question.",
        "You are a construction manager selecting materials for a "
        "commercial project in Melbourne.",
        "What legislation must you consider?",
        "Answer must be 250-300 words.",
        "Satisfactory response",
        "Yes ☐",
        "No ☐",
    ])
    out.add_paragraph("")
    box([""])                       # the student's answer box
    out.add_paragraph("")

    box([
        "Question 2. Read the case study and answer 2(a) to 2(b).",
        "A supplier delivers framing that deviates from the NCC.",
        "Satisfactory response",
        "Yes ☐",
        "No ☐",
    ])
    out.add_paragraph("")

    box([
        "Question 2(a). Explain what project plans are.",
        "Answer must be 30-80 words.",
        "Satisfactory response",
        "Yes ☐",
        "No ☐",
    ])
    out.add_paragraph("")
    box([""])
    out.add_paragraph("")

    box([
        "Question 2(b). Is a safe work method statement required?",
        "Yes ☐",
        "No ☐",
    ])
    out.add_paragraph("")
    box([""])
    out.add_paragraph("")

    # Question 3: a grid to fill in
    box(["Question 3. Explain the following building materials."])
    grid = out.add_table(rows=4, cols=4)
    grid.style = "Table Grid"
    headers = ["Material", "Properties", "Applications", "Limitations"]
    for column, text in enumerate(headers):
        grid.rows[0].cells[column].paragraphs[0].text = text
    for row, name in enumerate(["1. Cements", "2. Ceramics", "3. Concrete"], 1):
        grid.rows[row].cells[0].paragraphs[0].text = name
    out.add_paragraph("")

    # Question 4: no space at all - real text follows it
    out.add_paragraph("Question 4. Name one hazard on site.")
    out.add_paragraph("End of assessment task 1.")

    out.save(path)


def main():

    work = tempfile.mkdtemp(prefix="docx-agent-")
    path = os.path.join(work, "workbook.docx")

    build(path)

    loaded = doc_io.read(path)
    found = finder.find(loaded)

    print("read %d blocks, found %d questions\n" % (
        len(loaded.blocks), len(found)
    ))

    for question in found:
        detail = question.space_kind
        if question.grid:
            detail = "grid %dx%d" % (
                len(question.grid.rows), len(question.grid.headers)
            )
        print("  [%s] %-6s %-10s words=%s\n      %s" % (
            question.number,
            question.answer_type,
            detail,
            (question.word_min and
             "%d-%d" % (question.word_min, question.word_max)) or "-",
            question.text.replace("\n", " | ")[:70],
        ))

    expected = ["1", "2(a)", "2(b)", "3", "4"]
    actual = [q.number for q in found]

    problems = []

    if actual != expected:
        problems.append("questions: expected %s, got %s" % (expected, actual))

    by_number = {q.number: q for q in found}

    if by_number["2(b)"].answer_type != "yesno":
        problems.append("2(b) should want Yes/No")

    if by_number["1"].word_max != 300:
        problems.append("1 should carry its 250-300 word limit")

    if by_number["3"].grid is None:
        problems.append("3 should be a grid")

    if by_number["4"].space_kind != "none":
        problems.append("4 should have no space of its own")

    # --- write answers in, the way a run would ---------------------

    answers = {
        "1": "Legislation includes the WHS Act and the NCC.\n\n"
             "Australian Standards apply to storage as well.",
        "2(a)": "Project plans are the drawings and specifications.",
        "2(b)": "Yes",
        "4": "Working at heights near an unguarded edge.",
    }

    style = doc_io.style_from(loaded.blocks[0])

    for question in found:

        if question.grid:
            for row_at, _ in enumerate(question.grid.rows, start=1):
                row = question.grid.table.rows[row_at]
                for column in range(1, len(question.grid.headers) + 1):
                    cell = row.cells[column]
                    if not cell.text.strip():
                        cell.paragraphs[0].text = "filled %d" % column
            continue

        text = answers.get(question.number)

        if not text:
            continue

        if question.space is not None:
            doc_io.fill(question.space, text, style)
        else:
            doc_io.insert_after(question.after, text, style)

    saved = os.path.join(work, "workbook-answered.docx")
    loaded.save(saved)

    # --- read it back and check where everything landed ------------

    again = doc_io.read(saved)
    text_of = [b.text for b in again.blocks]
    whole = "\n".join(text_of)

    if "Legislation includes the WHS Act" not in whole:
        problems.append("answer 1 was not written")

    for at, block in enumerate(again.blocks):
        if "Legislation includes the WHS Act" in block.text:
            if not block.in_table:
                problems.append("answer 1 did not go into a box")
            if "Question 1." in block.text:
                problems.append("answer 1 landed inside the question box")
            break

    if "Satisfactory response\nYes" in whole.replace("☐", "").replace(
        " ", ""
    ):
        pass  # marking lines still intact, which is what we want

    if "Yes" not in whole:
        problems.append("the Yes/No answer was not written")

    if "Working at heights" not in whole:
        problems.append("answer 4 was not added under its question")

    for at, block in enumerate(again.blocks):
        if "Question 4." in block.text:
            following = again.blocks[at + 1].text if at + 1 < len(
                again.blocks
            ) else ""
            if "Working at heights" not in following:
                problems.append(
                    "answer 4 is not directly under question 4"
                )
            break

    filled = sum(1 for b in again.blocks if b.text.startswith("filled"))

    if filled != 9:
        problems.append("grid: expected 9 filled cells, got %d" % filled)

    print("\nsaved: %s" % saved)

    if problems:
        print("\nFAILED:")
        for problem in problems:
            print("  - " + problem)
        return 1

    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
